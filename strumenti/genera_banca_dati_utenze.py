# =========================================================================
# genera_banca_dati_utenze.py — modello Excel della Banca Dati Utenze CWP.
#
# Scopo: genera public/modelli/Banca_dati_utenze_CWP_v1.xlsx, il file che
# il Comune o il gestore compila OGNI ANNO con l'elenco COMPLETO delle
# utenze (richiesta di Luca Moretti del 22/10/2026: il calcolo dei crediti
# è per utenza e le utenze cambiano di continuo, quindi ogni anno si
# ricarica tutto, domestiche e non domestiche).
#
# Fogli: Istruzioni (spiegazione di ogni colonna ed esempio), Utenze_domestiche,
# Utenze_non_domestiche, Categorie (categorie DPR 158/1999, solo riferimento).
#
# Colonne: una riga per utenza; per ogni frazione i kg dell'anno e, a
# parte, gli svuotamenti. Gli id delle frazioni sono quelli di
# src/lib/crediti/parametri.js e della funzione SQL cwp_tipologie_rifiuto()
# (script 40); la tabella del database è calcolo_utenza. Se si cambiano
# colonne o frazioni qui, cambiarle anche nel lettore del portale e nei
# documenti (docs/crediti-cwp.md, manuale amministratore, pagina Formato dati).
#
# Uso: python3 strumenti/genera_banca_dati_utenze.py (serve openpyxl),
# poi commit del file .xlsx in public/modelli/.
#
# Registro sviluppo e deploy: capitolo 15 «Calcolo crediti».
# =========================================================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import os

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "public", "modelli", "Banca_dati_utenze_CWP_v1.xlsx")
VERDE = "2F5D3A"; VERDE_CHIARO = "E4EDE1"; GIALLO = "FFF4C2"; GRIGIO = "5A5E54"; ROSSO = "A23B2C"; OCRA = "B06A2F"
F = "Arial"
f_base = Font(name=F, size=10)
f_tit = Font(name=F, size=15, bold=True, color=VERDE)
f_sez = Font(name=F, size=12, bold=True, color=VERDE)
f_sub = Font(name=F, size=10, italic=True, color=GRIGIO)
f_h = Font(name=F, size=9.5, bold=True, color="FFFFFF")
f_b = Font(name=F, size=10, bold=True)
fill_h = PatternFill("solid", fgColor=VERDE)
fill_obbl = PatternFill("solid", fgColor=OCRA)
fill_soft = PatternFill("solid", fgColor=VERDE_CHIARO)
thin = Side(style="thin", color="D9D4C5")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

# Frazioni: (id del portale, nome nella colonna). L'indifferenziato è
# obbligatorio; l'ordine è quello delle colonne.
FRAZIONI = [
    ("indifferenziato", "Indifferenziato"),
    ("organico", "Organico"),
    ("carta", "Carta e cartone"),
    ("plastica", "Plastica"),
    ("vetro", "Vetro"),
    ("metallo", "Metallo"),
    ("verde", "Verde"),
    ("legno", "Legno"),
    ("tessili", "Tessili"),
    ("raee", "RAEE"),
    ("ingombranti", "Ingombranti"),
    ("inerti", "Inerti"),
]

# Colonne anagrafiche: (intestazione, larghezza, obbligatoria, tipo, spiegazione).
# tipo: "testo", "intero" (>= 1), "decimale" (> 0).
ANAG_UD = [
    ("Codice utenza", 14, True, "testo", "Codice dell'utenza nel sistema del Comune o del gestore (es. codice TARI). Unico nel file: due righe non possono avere lo stesso codice."),
    ("Nome capofamiglia", 26, True, "testo", "Nome e cognome dell'intestatario dell'utenza domestica."),
    ("Indirizzo", 32, True, "testo", "Indirizzo dell'utenza: via, numero civico ed eventuale interno."),
    ("Categoria tariffaria", 12, False, "testo", "Classe della quota variabile nella delibera TARI del Comune (di solito il numero di componenti: 1, 2, … 6 e oltre). Se presente, il portale la propone come categoria di simulazione, modificabile."),
    ("N. componenti", 11, False, "intero", "Numero di componenti del nucleo familiare (facoltativo). Serve a collocare l'utenza nella classe della delibera TARI; se manca si usano i valori medi."),
    ("Mq abitazione", 11, False, "decimale", "Superficie dell'abitazione in metri quadrati (facoltativa)."),
]
ANAG_UND = [
    ("Codice utenza", 14, True, "testo", "Codice dell'utenza nel sistema del Comune o del gestore (es. codice TARI). Unico nel file."),
    ("Ragione sociale", 28, True, "testo", "Denominazione dell'attività (ditta, società, ente)."),
    ("Indirizzo", 32, True, "testo", "Indirizzo dei locali dell'utenza."),
    ("Categoria tariffaria", 12, False, "testo", "Numero della categoria della delibera TARI del Comune (DPR 158/1999, es. 22 per ristoranti; elenco di riferimento nel foglio Categorie). Se presente, il portale la propone come categoria di simulazione; se manca, la Fondazione la assegna nel portale. Compilatela sempre se la conoscete."),
    ("Mq locali", 11, False, "decimale", "Superficie tassata dei locali in metri quadrati (facoltativa; se manca si usano i valori medi della categoria)."),
]
TOTALE = ("Produzione totale (kg)", 13, True, "decimale0", "Rifiuti totali prodotti dall'utenza nell'anno, in kg (somma di tutte le frazioni). Il portale la confronta con la somma delle colonne delle frazioni.")

# Nome in minuscolo dentro le frasi, tranne le sigle (RAEE).
basso = lambda n: n if n.isupper() else n.lower()

def colonne(anag):
    cols = list(anag) + [TOTALE]
    for fid, nome in FRAZIONI:
        obbl = fid == "indifferenziato"
        cols.append((f"{nome} (kg)", 12, obbl, "decimale0",
                     ("Kg di indifferenziato (secco residuo) prodotti nell'anno. Obbligatorio: anche 0." if obbl
                      else f"Kg di {basso(nome)} prodotti nell'anno. Vuoto = frazione non raccolta o non misurata per l'utenza.")))
    for fid, nome in FRAZIONI:
        cols.append((f"Svuotamenti {basso(nome)}", 12, False, "intero0",
                     f"Numero di svuotamenti (o conferimenti) di {basso(nome)} nell'anno. Facoltativo."))
    return cols

COL_UD = colonne(ANAG_UD)
COL_UND = colonne(ANAG_UND)

ES_UD = ["UD000123", "Mario Rossi", "Via Roma 12, int. 3", "3", 3, 95, 630, 210, 180, 110, 50, 70, 10, None, None, None, None, None, None,
         26, 52, 13, 26, 13, None, None, None, None, None, None, None]
ES_UND = ["UND000045", "Bar Centrale s.r.l.", "Piazza Duomo 2", "24", 80, 4200, 1300, 900, 600, 400, 1000, None, None, None, None, None, None, None,
          104, 156, 52, 52, 52, None, None, None, None, None, None, None]

# Categorie per le utenze non domestiche, DPR 158/1999 (Comuni con più di
# 5.000 abitanti). Solo riferimento: fa fede la delibera del Comune.
CATEGORIE = [
    (1, "Musei, biblioteche, scuole, associazioni, luoghi di culto"),
    (2, "Cinematografi e teatri"),
    (3, "Autorimesse e magazzini senza alcuna vendita diretta"),
    (4, "Campeggi, distributori carburanti, impianti sportivi"),
    (5, "Stabilimenti balneari"),
    (6, "Esposizioni, autosaloni"),
    (7, "Alberghi con ristorante"),
    (8, "Alberghi senza ristorante"),
    (9, "Case di cura e riposo"),
    (10, "Ospedali"),
    (11, "Uffici, agenzie, studi professionali"),
    (12, "Banche ed istituti di credito"),
    (13, "Negozi abbigliamento, calzature, libreria, cartoleria, ferramenta e altri beni durevoli"),
    (14, "Edicola, farmacia, tabaccaio, plurilicenze"),
    (15, "Negozi particolari quali filatelia, tende e tessuti, tappeti, cappelli e ombrelli, antiquariato"),
    (16, "Banchi di mercato beni durevoli"),
    (17, "Attività artigianali tipo botteghe: parrucchiere, barbiere, estetista"),
    (18, "Attività artigianali tipo botteghe: falegname, idraulico, fabbro, elettricista"),
    (19, "Carrozzeria, autofficina, elettrauto"),
    (20, "Attività industriali con capannoni di produzione"),
    (21, "Attività artigianali di produzione beni specifici"),
    (22, "Ristoranti, trattorie, osterie, pizzerie, pub"),
    (23, "Mense, birrerie, amburgherie"),
    (24, "Bar, caffè, pasticceria"),
    (25, "Supermercato, pane e pasta, macelleria, salumi e formaggi, generi alimentari"),
    (26, "Plurilicenze alimentari e/o miste"),
    (27, "Ortofrutta, pescherie, fiori e piante, pizza al taglio"),
    (28, "Ipermercati di generi misti"),
    (29, "Banchi di mercato generi alimentari"),
    (30, "Discoteche, night club"),
]

wb = Workbook()
ist = wb.active; ist.title = "Istruzioni"

def foglio_dati(nome, cols, n_righe=200000):
    ws = wb.create_sheet(nome)
    for j, (h, w, obbl, tipo, _) in enumerate(cols, start=1):
        c = ws.cell(1, j, h); c.font = f_h; c.fill = fill_obbl if obbl else fill_h; c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
        lettera = get_column_letter(j)
        if tipo in ("intero", "intero0"):
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1" if tipo == "intero" else "0", allow_blank=True,
                                showErrorMessage=True, errorTitle="Valore non ammesso", error="Serve un numero intero senza decimali" + (" (almeno 1)." if tipo == "intero" else " (0 o più)."))
        elif tipo in ("decimale", "decimale0"):
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                                showErrorMessage=True, errorTitle="Valore non ammesso", error="Serve un numero (0 o più), senza unità di misura.")
        else:
            dv = None
        if dv:
            ws.add_data_validation(dv); dv.add(f"{lettera}2:{lettera}{n_righe}")
            for r in range(2, 502):
                ws.cell(r, j).number_format = "#,##0" if tipo in ("intero", "intero0") else "#,##0.0"
        else:
            for r in range(2, 502):
                ws.cell(r, j).number_format = "@"
    ws.row_dimensions[1].height = 54
    ws.freeze_panes = "B2"
    return ws

foglio_dati("Utenze_domestiche", COL_UD)
foglio_dati("Utenze_non_domestiche", COL_UND)

# ---------------- Categorie ----------------
cat = wb.create_sheet("Categorie")
cat["A1"] = "Categorie delle utenze non domestiche (DPR 158/1999)"; cat["A1"].font = f_tit
cat["A2"] = "Elenco di riferimento per i Comuni con più di 5.000 abitanti. Nella colonna Categoria scrivete il numero della categoria della VOSTRA delibera TARI, che fa fede (i Comuni più piccoli hanno 21 categorie con numerazione diversa)."; cat["A2"].font = f_sub
cat["A2"].alignment = wrap; cat.merge_cells("A2:B2"); cat.row_dimensions[2].height = 42
for j, h in enumerate(["Categoria", "Attività"], start=1):
    c = cat.cell(4, j, h); c.font = f_h; c.fill = fill_h; c.border = box
for i, (n, d) in enumerate(CATEGORIE):
    cat.cell(5 + i, 1, n).border = box; c = cat.cell(5 + i, 2, d); c.border = box; c.font = f_base
cat.column_dimensions["A"].width = 12; cat.column_dimensions["B"].width = 90

# ---------------- Istruzioni ----------------
ist.column_dimensions["A"].width = 3
ist.column_dimensions["B"].width = 30; ist.column_dimensions["C"].width = 14; ist.column_dimensions["D"].width = 90
r = 1
def riga(testo, font=f_base, fill=None, altezza=None):
    global r
    c = ist.cell(r, 2, testo); c.font = font; c.alignment = wrap
    ist.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    if fill:
        for cc in range(2, 5): ist.cell(r, cc).fill = fill
    if altezza: ist.row_dimensions[r].height = altezza
    r += 1

riga("Programma Carbon WastePrint® — Banca Dati Utenze per il calcolo dei crediti", f_tit, altezza=24)
riga("Versione 1.0 · ottobre 2026 · Fondazione Operate", f_sub)
r += 1
riga("A cosa serve", f_sez)
riga("Il calcolo dei crediti di carbonio è fatto utenza per utenza. Ogni anno il Comune o il gestore del servizio di raccolta compila questo file con l'ELENCO COMPLETO delle utenze dell'anno di riferimento, domestiche e non domestiche, e lo carica nel portale nel calcolo dell'anno. L'elenco si ricarica sempre per intero, perché le utenze cambiano di continuo (nuove utenze, cessazioni, traslochi, cambi di categoria).", altezza=58)
r += 1
riga("Come compilarlo", f_sez)
for p in [
    "1. Foglio Utenze_domestiche: una riga per ogni utenza domestica attiva nell'anno (anche se attiva solo per una parte dell'anno).",
    "2. Foglio Utenze_non_domestiche: una riga per ogni utenza non domestica (attività, enti, uffici).",
    "3. Partite dalla riga 2, sotto le intestazioni. Non cambiate, spostate o cancellate le intestazioni: il portale le usa per riconoscere le colonne.",
    "4. Le colonne con l'intestazione arancione sono obbligatorie; quelle verdi sono facoltative e si possono lasciare vuote.",
    "5. Numeri senza unità di misura e senza testo; i decimali possono avere la virgola. Kg e svuotamenti sono i totali dell'anno di riferimento.",
    "6. Per ogni utenza: produzione totale in kg, kg di indifferenziato (obbligatorio, anche 0) e kg delle altre frazioni raccolte. Una frazione non raccolta o non misurata per l'utenza si lascia vuota (vuoto è diverso da 0).",
    "7. La produzione totale deve essere uguale alla somma delle frazioni (il portale segnala le differenze oltre l'1%).",
    "8. Categoria tariffaria: indicate la classe (domestiche) o la categoria (non domestiche) della vostra delibera TARI. Nel portale ogni utenza ha anche una «categoria di simulazione», quella usata dal calcolo per i rifiuti attesi: il portale la propone uguale alla tariffaria e la Fondazione può cambiarla.",
    "9. Oltre 1 milione di righe o 50 MB: salvate ciascun foglio come CSV (separatore punto e virgola) con le stesse intestazioni e lo stesso nome del foglio (es. Utenze_domestiche.csv) e riunite i file in un archivio ZIP.",
]:
    riga(p, altezza=30)
r += 1
riga("Dati personali", Font(name=F, size=12, bold=True, color=ROSSO))
riga("Questo file contiene dati personali (nomi e indirizzi). Va trasmesso SOLO caricandolo nel portale Carbon WastePrint, mai per email. Il Comune è titolare del trattamento; la Fondazione Operate tratta i dati per suo conto solo per il calcolo dei crediti, secondo l'accordo sul trattamento dei dati. Nel portale i dati li vedono solo gli Operatori e gli Amministratori della Fondazione, ogni consultazione è registrata e nomi e indirizzi non compaiono nei rapporti né nelle pagine pubbliche. Non inserite altri dati personali (codice fiscale, telefono, email).", Font(name=F, size=10, color=ROSSO), fill=PatternFill("solid", fgColor="F6DCD6"), altezza=84)
r += 1

def spiegazione(titolo, cols, esempio):
    global r
    riga(titolo, f_b)
    for j, h in enumerate(["Colonna", "Obbligatoria", "Cosa contiene"]):
        c = ist.cell(r, 2 + j, h); c.font = f_h; c.fill = fill_h; c.border = box
    r += 1
    # Le 12 colonne degli svuotamenti si spiegano in una riga sola.
    sv = [c for c in cols if c[0].startswith("Svuotamenti")]
    for h, w, obbl, tipo, nota in [c for c in cols if not c[0].startswith("Svuotamenti")]:
        ist.cell(r, 2, h).font = f_b; ist.cell(r, 3, "sì" if obbl else "no").font = f_base
        ist.cell(r, 4, nota).font = f_base
        for cc in range(2, 5): ist.cell(r, cc).border = box; ist.cell(r, cc).alignment = wrap
        ist.row_dimensions[r].height = 30 if len(nota) > 95 else 16
        r += 1
    ist.cell(r, 2, f"Svuotamenti … (12 colonne: {', '.join(basso(n) for _, n in FRAZIONI)})").font = f_b
    ist.cell(r, 3, "no").font = f_base
    ist.cell(r, 4, "Numero di svuotamenti del contenitore (o di conferimenti) di quella frazione nell'anno. Intero, 0 o più; vuoto se non rilevato.").font = f_base
    for cc in range(2, 5): ist.cell(r, cc).border = box; ist.cell(r, cc).alignment = wrap
    ist.row_dimensions[r].height = 30
    r += 1
    riga("Esempio di riga (solo qui: i fogli da compilare partono vuoti): " + " · ".join(
        f"{cols[i][0]} = {v}" for i, v in enumerate(esempio) if v is not None), Font(name=F, size=9.5, italic=True), fill=fill_soft, altezza=58)
    r += 1

riga("Le colonne", f_sez)
spiegazione("Foglio Utenze_domestiche", COL_UD, ES_UD)
spiegazione("Foglio Utenze_non_domestiche", COL_UND, ES_UND)
riga("Legenda: intestazione arancione = obbligatoria; intestazione verde = facoltativa.", f_sub)
riga("Per domande: Fondazione Operate, tramite la pratica nel portale Carbon WastePrint.", f_sub)
ist.sheet_view.showGridLines = False

for ws in wb.worksheets:
    ws.sheet_properties.tabColor = VERDE if ws.title == "Istruzioni" else (OCRA if ws.title == "Categorie" else "8AA88F")
    ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.active = 0
wb.save(OUT)
print("ok", OUT, len(COL_UD), "colonne domestiche,", len(COL_UND), "non domestiche")
