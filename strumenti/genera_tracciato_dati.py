# Genera public/modelli/Tracciato_dati_CWP_v1.xlsx (tracciato dati standard CWP).
# Uso: python3 strumenti/genera_tracciato_dati.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "public", "modelli", "Tracciato_dati_CWP_v1.xlsx")
VERDE = "2F5D3A"; VERDE_CHIARO = "E4EDE1"; CARTA = "F6F4EE"; GIALLO = "FFF4C2"; GRIGIO = "5A5E54"; ROSSO = "A23B2C"
F = "Arial"
f_base = Font(name=F, size=10)
f_tit = Font(name=F, size=15, bold=True, color=VERDE)
f_sub = Font(name=F, size=10, italic=True, color=GRIGIO)
f_h = Font(name=F, size=10, bold=True, color="FFFFFF")
f_b = Font(name=F, size=10, bold=True)
f_es = Font(name=F, size=9.5, italic=True, color=GRIGIO)
fill_h = PatternFill("solid", fgColor=VERDE)
fill_in = PatternFill("solid", fgColor=GIALLO)
fill_soft = PatternFill("solid", fgColor=VERDE_CHIARO)
fill_carta = PatternFill("solid", fgColor=CARTA)
thin = Side(style="thin", color="D9D4C5")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")

TIPOLOGIE = [
  ("indifferenziato", "Indifferenziato (secco residuo)", "200301"),
  ("organico", "Frazione organica (umido)", "200108, 200302"),
  ("verde", "Scarto verde (sfalci e potature)", "200201"),
  ("carta", "Carta e cartone", "200101, 150101"),
  ("plastica", "Plastica (imballaggi)", "150102, 200139"),
  ("vetro", "Vetro", "150107, 200102"),
  ("metallo", "Metallo", "150104, 200140"),
  ("legno", "Legno", "150103, 200138"),
  ("tessili", "Tessili", "200110, 200111"),
  ("raee", "RAEE", "200121, 200123, 200135, 200136"),
  ("ingombranti", "Ingombranti", "200307"),
  ("inerti", "Rifiuti inerti", "170107, 170904"),
]
BENI = [("maglietta","Maglietta"),("maglione","Maglione"),("vestito","Vestito generico"),("mobile","Mobile generico"),("aee","AEE generici (apparecchiature elettriche)")]
PUNTI = [("domiciliare","Raccolta domiciliare (porta a porta)"),("struttura","Struttura informatizzata (cassonetto con accesso controllato)"),("centro_raccolta","Centro di raccolta")]

# Fogli di dati: (nome, titolo, [(intestazione, larghezza, obbligatorio, formato, lista, nota)], esempio)
FOGLI = [
 ("2_Rifiuti", "Rifiuti raccolti nel territorio comunale",
  "Una riga per codice EER, con le quantità dell'anno di riferimento (Catasto ISPRA, MUD o O.R.So.). Più codici EER possono avere la stessa Tipologia.",
  [("Codice EER",12,True,"@",None,"Codice europeo del rifiuto, 6 cifre senza spazi (es. 200301)"),
   ("Descrizione",40,False,None,None,"Descrizione del codice EER"),
   ("Tipologia",18,True,None,"tipologie","Codice della tipologia CWP (foglio Codici)"),
   ("Quantità (t/anno)",16,True,"#,##0.000",None,"Tonnellate raccolte nell'anno, con decimali"),
   ("Solo domestiche",14,False,None,"sino","sì se prodotto solo da utenze domestiche (es. verde, inerti)")],
  ["200301","Rifiuti urbani non differenziati","indifferenziato",1250.4,"no"]),
 ("3_Classi_domestiche", "Classi delle utenze domestiche (delibera TARI)",
  "Una riga per classe della quota variabile, con il numero di utenze e il coefficiente Kb della delibera.",
  [("Classe",22,True,None,None,"Nome della classe come in delibera (es. 1 componente)"),
   ("N. utenze",12,True,"#,##0",None,"Numero di utenze della classe"),
   ("Kb",10,True,"0.000",None,"Coefficiente Kb della quota variabile")],
  ["1 componente",1200,0.6]),
 ("4_Classi_non_domestiche", "Categorie delle utenze non domestiche (delibera TARI)",
  "Una riga per categoria, con numero di utenze, superficie complessiva e coefficiente Kd della delibera.",
  [("Categoria",40,True,None,None,"Categoria come in delibera (es. 1 - Musei, biblioteche, scuole)"),
   ("N. utenze",12,True,"#,##0",None,"Numero di utenze della categoria"),
   ("Superficie (m²)",16,True,"#,##0",None,"Superficie tassata complessiva della categoria"),
   ("Kd (kg/m²/anno)",16,True,"0.000",None,"Coefficiente Kd della quota variabile")],
  ["1 - Musei, biblioteche, scuole",12,5400,2.6]),
 ("5_Conferimenti", "Banca Dati Utenze — conferimenti annui (dati riassunti)",
  "Una riga per utenza e tipologia di rifiuto con i TOTALI dell'anno. Se il vostro sistema esporta le singole letture, compilate invece il foglio 6_Letture.",
  [("Utenza",14,True,None,None,"Codice anonimo dell'utenza, stabile negli anni. NIENTE nomi, indirizzi o codici fiscali"),
   ("Macrocategoria",15,True,None,"macro","UD = domestica, UnD = non domestica"),
   ("Classe",26,True,None,None,"Classe o categoria della delibera (identica ai fogli 3 e 4)"),
   ("Tipologia rifiuto",18,True,None,"tipologie","Codice della tipologia CWP (foglio Codici)"),
   ("Svuotamenti",13,False,"#,##0",None,"Numero di svuotamenti del contenitore nell'anno (porta a porta)"),
   ("Volume contenitore (L)",14,False,"#,##0",None,"Litri del contenitore svuotato"),
   ("Peso strutture informatizzate (kg)",16,False,"#,##0.0",None,"Kg conferiti nell'anno a strutture con pesatura"),
   ("Volume strutture informatizzate (L)",16,False,"#,##0",None,"Litri conferiti nell'anno a strutture senza pesatura"),
   ("Peso centri di raccolta (kg)",15,False,"#,##0.0",None,"Kg conferiti nell'anno ai centri di raccolta"),
   ("Volume centri di raccolta (L)",15,False,"#,##0",None,"Litri conferiti nell'anno ai centri di raccolta"),
   ("N. utenze (valori medi)",14,False,"#,##0",None,"Solo per stime ex ante: quante utenze rappresenta la riga. Altrimenti vuoto")],
  ["U000123","UD","1 componente","indifferenziato",18,40,None,None,None,None,None]),
 ("6_Letture", "Banca Dati Utenze — singole letture (dati grezzi)",
  "In alternativa al foglio 5: una riga per ogni svuotamento o conferimento letto dal sistema. Il portale calcola da solo i totali annui. Oltre 1 milione di righe: usate file CSV dentro uno ZIP.",
  [("Utenza",14,True,None,None,"Codice anonimo dell'utenza, stabile negli anni"),
   ("Macrocategoria",15,True,None,"macro","UD = domestica, UnD = non domestica"),
   ("Classe",26,True,None,None,"Classe o categoria della delibera"),
   ("Tipologia rifiuto",18,True,None,"tipologie","Codice della tipologia CWP"),
   ("Data",12,True,"DD/MM/YYYY",None,"Data del conferimento (gg/mm/aaaa)"),
   ("Punto di conferimento",22,True,None,"punti","domiciliare, struttura o centro_raccolta"),
   ("Volume contenitore (L)",14,False,"#,##0",None,"Litri del contenitore svuotato o conferito"),
   ("Peso (kg)",11,False,"#,##0.0",None,"Kg, se il conferimento è pesato")],
  ["U000123","UD","1 componente","indifferenziato","15/01/2026","domiciliare",40,None]),
 ("7_Riuso", "Beni conferiti a riuso",
  "Una riga per utenza e tipo di bene conferito al centro del riuso nell'anno.",
  [("Utenza",14,True,None,None,"Codice anonimo dell'utenza (lo stesso dei fogli 5 o 6)"),
   ("Macrocategoria",15,True,None,"macro","UD o UnD"),
   ("Classe",26,True,None,None,"Classe o categoria della delibera"),
   ("Bene",14,True,None,"beni","Codice del bene (foglio Codici)"),
   ("Pezzi",9,False,"#,##0",None,"Numero di pezzi"),
   ("Peso (kg)",11,False,"#,##0.0",None,"Peso complessivo, se misurato"),
   ("N. utenze (valori medi)",14,False,"#,##0",None,"Solo per stime ex ante")],
  ["U000123","UD","1 componente","maglione",3,None,None]),
]

wb = Workbook()
ist = wb.active; ist.title = "Istruzioni"

# ---------------- Codici ----------------
cod = wb.create_sheet("Codici")
cod["A1"] = "Codici ammessi"; cod["A1"].font = f_tit
cod["A2"] = "Usate esattamente questi codici nelle colonne con il menu a tendina."; cod["A2"].font = f_sub
def tabella_codici(r0, c0, titolo, intest, righe):
    cod.cell(r0, c0, titolo).font = f_b
    for j, h in enumerate(intest):
        c = cod.cell(r0+1, c0+j, h); c.font = f_h; c.fill = fill_h; c.border = box
    for i, riga in enumerate(righe):
        for j, v in enumerate(riga):
            c = cod.cell(r0+2+i, c0+j, v); c.font = f_base; c.border = box
    return (r0+2, r0+1+len(righe))
r_tip = tabella_codici(4, 1, "Tipologie di rifiuto", ["Codice", "Descrizione", "Codici EER tipici"], TIPOLOGIE)
r_mac = tabella_codici(4, 5, "Macrocategorie", ["Codice", "Descrizione"], [("UD","Utenza domestica"),("UnD","Utenza non domestica")])
r_sn  = tabella_codici(9, 5, "Sì / No", ["Valore"], [("sì",),("no",)])
r_ben = tabella_codici(19, 1, "Beni a riuso (Metodologia, Tabella 5)", ["Codice", "Descrizione"], BENI)
r_pun = tabella_codici(19, 5, "Punti di conferimento", ["Codice", "Descrizione"], PUNTI)
for col, w in zip("ABCDEFG", [18, 36, 30, 3, 18, 58, 3]): cod.column_dimensions[col].width = w

liste = {
 "tipologie": f"Codici!$A${r_tip[0]}:$A${r_tip[1]}",
 "macro": f"Codici!$E${r_mac[0]}:$E${r_mac[1]}",
 "sino": f"Codici!$E${r_sn[0]}:$E${r_sn[1]}",
 "beni": f"Codici!$A${r_ben[0]}:$A${r_ben[1]}",
 "punti": f"Codici!$E${r_pun[0]}:$E${r_pun[1]}",
}

# ---------------- 1_Anagrafica ----------------
an = wb.create_sheet("1_Anagrafica", 1)
an["A1"] = "Anagrafica e informazioni generali"; an["A1"].font = f_tit
an["A2"] = "Compilate le celle gialle."; an["A2"].font = f_sub
campi = [
 ("Ente (Comune o Unione)", "Denominazione ufficiale", None),
 ("Codice ISTAT", "6 cifre", "@"),
 ("Provincia", "Sigla", None),
 ("Anno di riferimento dei dati", "Es. 2025", "@"),
 ("Gestore del servizio di raccolta", "Ragione sociale", None),
 ("Software di gestione dei dati", "Nome del programma da cui sono estratti i dati", None),
 ("Tariffa puntuale attiva", "sì / no", "sino"),
 ("Data di avvio della tariffa puntuale", "gg/mm/aaaa, se attiva", "DD/MM/YYYY"),
 ("PEF % utenze domestiche", "Dalla delibera tariffaria (es. 65)", "0.00"),
 ("PEF % utenze non domestiche", "Dalla delibera tariffaria (es. 35)", "0.00"),
 ("Estremi della delibera tariffaria", "Numero e data", None),
 ("Referente per i dati (ufficio / ruolo)", "Es. Ufficio Ambiente", None),
 ("Email di contatto dell'ufficio", "Indirizzo istituzionale", None),
 ("Data di compilazione", "gg/mm/aaaa", "DD/MM/YYYY"),
]
for i, h in enumerate(["Voce", "Valore", "Indicazioni"]):
    c = an.cell(4, 1+i, h); c.font = f_h; c.fill = fill_h; c.border = box
dv_an = DataValidation(type="list", formula1=liste["sino"], allow_blank=True)
an.add_data_validation(dv_an)
for i, (voce, nota, fmt) in enumerate(campi):
    r = 5+i
    a = an.cell(r, 1, voce); a.font = f_b; a.border = box
    b = an.cell(r, 2); b.fill = fill_in; b.border = box; b.font = f_base
    if fmt == "sino": dv_an.add(b)
    elif fmt: b.number_format = fmt
    c = an.cell(r, 3, nota); c.font = f_es; c.border = box
pef_ud, pef_und = 5+8, 5+9
r_chk = 5+len(campi)+1
an.cell(r_chk, 1, "Controllo PEF (deve essere 100)").font = f_es
k = an.cell(r_chk, 2, f'=IF(AND(B{pef_ud}="",B{pef_und}=""),"",B{pef_ud}+B{pef_und})'); k.font = f_base; k.number_format = "0.00"
an.column_dimensions["A"].width = 40; an.column_dimensions["B"].width = 34; an.column_dimensions["C"].width = 48

# ---------------- Fogli dati ----------------
for pos, (nome, titolo, descr, colonne, esempio) in enumerate(FOGLI, start=2):
    ws = wb.create_sheet(nome, pos)
    for j, (h, w, obbl, fmt, lista, nota) in enumerate(colonne, start=1):
        c = ws.cell(1, j, h); c.font = f_h; c.fill = fill_h; c.border = box
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
        n_righe = 200000 if nome == "6_Letture" else 20000
        if fmt:
            for r in range(2, 302): ws.cell(r, j).number_format = fmt
        if lista:
            dv = DataValidation(type="list", formula1=liste[lista], allow_blank=True,
                                showErrorMessage=True, errorTitle="Valore non ammesso",
                                error="Scegliete un codice dal menu (foglio Codici).")
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(j)}2:{get_column_letter(j)}{n_righe}")
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "A2"

# ---------------- Istruzioni ----------------
ist.column_dimensions["A"].width = 3
for col, w in zip("BCDEFGHIJKL", [26, 18, 18, 18, 16, 16, 16, 16, 16, 16, 16]):
    ist.column_dimensions[col].width = w
r = 1
def riga(testo, font=f_base, fill=None, altezza=None, unisci=True):
    global r
    c = ist.cell(r, 2, testo); c.font = font; c.alignment = wrap
    if unisci: ist.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
    if fill:
        for cc in range(2, 13): ist.cell(r, cc).fill = fill
    if altezza: ist.row_dimensions[r].height = altezza
    r += 1

riga("Programma Carbon WastePrint® — Tracciato dati per il calcolo dei crediti", f_tit, altezza=24)
riga("Versione 1.0 · ottobre 2026 · Fondazione Operate", f_sub)
r += 1
riga("A cosa serve", Font(name=F, size=12, bold=True, color=VERDE))
riga("Questo file raccoglie i dati che il Comune e il gestore del servizio di raccolta forniscono alla Fondazione Operate per calcolare i crediti di carbonio secondo la Metodologia Carbon WastePrint v1.2. Tutti gli enti usano lo stesso formato: così i dati si caricano nel portale senza rielaborazioni.", altezza=44)
r += 1
riga("Come compilarlo", Font(name=F, size=12, bold=True, color=VERDE))
passi = [
 "1. Compilate il foglio 1_Anagrafica (celle gialle).",
 "2. Compilate i fogli 2, 3, 4 e 7 partendo dalla riga 2, sotto le intestazioni. Non cambiate né spostate le intestazioni.",
 "3. Per la Banca Dati Utenze compilate UNO dei due fogli: 5_Conferimenti (totali annui per utenza) oppure 6_Letture (singoli svuotamenti, se il vostro sistema esporta questi).",
 "4. Nelle colonne con il menu a tendina usate solo i codici del foglio Codici.",
 "5. Numeri senza unità di misura e senza testo; i decimali possono avere la virgola. Date nel formato gg/mm/aaaa.",
 "6. Un foglio non pertinente si lascia vuoto (ad esempio 7_Riuso se non c'è un centro del riuso).",
 "7. Caricate il file nel portale, nella pratica del vostro progetto. Se è troppo grande (oltre 50 MB o oltre 1 milione di righe) salvate ogni foglio come CSV con le stesse intestazioni (es. 6_Letture.csv) e riunite i file in un archivio ZIP.",
]
for p in passi: riga(p, altezza=28)
r += 1
riga("Protezione dei dati personali — obbligatorio", Font(name=F, size=12, bold=True, color=ROSSO))
riga("Il file NON deve contenere dati che identificano le persone: niente nomi e cognomi, indirizzi, codici fiscali, partite IVA, email o telefoni delle utenze. Ogni utenza si indica con un CODICE anonimo (ad esempio il codice interno del gestore o un codice progressivo), sempre lo stesso negli anni, in modo da confrontare i dati senza sapere a chi appartengono. La tabella di corrispondenza tra codici e persone resta al Comune o al gestore e non va inviata.", Font(name=F, size=10, color=ROSSO), fill=PatternFill("solid", fgColor="F6DCD6"), altezza=70)
r += 1
riga("Contenuto dei fogli ed esempio di compilazione", Font(name=F, size=12, bold=True, color=VERDE))
riga("Per ogni foglio: le colonne (★ = obbligatoria), cosa contengono e una riga di esempio con il formato atteso. L'esempio è solo qui: i fogli da compilare partono vuoti.", f_sub, altezza=28)
for nome, titolo, descr, colonne, esempio in FOGLI:
    r += 1
    riga(f"{nome} — {titolo}", f_b)
    riga(descr, altezza=30)
    for j, (h, w, obbl, fmt, lista, nota) in enumerate(colonne):
        c = ist.cell(r, 2+j, ("★ " if obbl else "") + h); c.font = Font(name=F, size=9, bold=True, color="FFFFFF"); c.fill = fill_h; c.border = box; c.alignment = Alignment(wrap_text=True, vertical="center")
        n = ist.cell(r+1, 2+j, nota); n.font = Font(name=F, size=8.5, color=GRIGIO); n.border = box; n.alignment = wrap
        e = ist.cell(r+2, 2+j, esempio[j]); e.font = Font(name=F, size=9.5, italic=True); e.fill = fill_soft; e.border = box; e.alignment = wrap
        if fmt and fmt not in ("@", "DD/MM/YYYY") and isinstance(esempio[j], (int, float)): e.number_format = fmt
    ist.row_dimensions[r].height = 40; ist.row_dimensions[r+1].height = 66
    ist.cell(r+2, 1, "es.").font = f_es
    r += 3
r += 1
riga("Legenda", Font(name=F, size=12, bold=True, color=VERDE))
c = ist.cell(r, 2, "Cella da compilare"); c.fill = fill_in; c.font = f_base; c.border = box
ist.cell(r, 3, "celle gialle del foglio 1_Anagrafica").font = f_es; r += 1
c = ist.cell(r, 2, "Intestazione"); c.fill = fill_h; c.font = f_h
ist.cell(r, 3, "non modificare: serve al portale per riconoscere le colonne").font = f_es; r += 1
c = ist.cell(r, 2, "Esempio"); c.fill = fill_soft; c.font = Font(name=F, size=9.5, italic=True)
ist.cell(r, 3, "riga dimostrativa, solo in questo foglio").font = f_es; r += 2
riga("Per domande sul tracciato: Fondazione Operate, tramite la pratica nel portale Carbon WastePrint.", f_sub)
ist.sheet_view.showGridLines = False
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = VERDE if ws.title in ("Istruzioni",) else ("B06A2F" if ws.title == "Codici" else "8AA88F")
for ws in wb.worksheets:
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
ist.print_area = f"A1:L{r}"
wb.active = 0
wb.save(OUT)
print("ok", OUT)
